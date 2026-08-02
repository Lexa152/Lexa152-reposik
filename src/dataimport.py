import csv
from typing import Any, Dict, List
from openpyxl import load_workbook


def import_data_csv(file_path: str) -> List[Dict[str, Any]]:
    """ Читает данные из CSV-файла и возвращает список словарей """
    result = []
    try:
        with open(file_path, mode='r', encoding='utf-8', newline='') as file:
            reader = csv.DictReader(file, delimiter=';')
            for i, row in enumerate(reader):
                result.append({
                    'id': row['id'],
                    'state': row['state'],
                    'date': row['date'],
                    'amount': row['amount'],
                    'currency_name': row['currency_name'],
                    'currency_code': row['currency_code'],
                    'from': row['from'],
                    'to': row['to'],
                    'description': row['description']
                })
    except FileNotFoundError:
        print(f"Файл не найден: {file_path}")
    except Exception as e:
        print(f"Произошла ошибка при чтении файла: {e}")
    return result


def import_data_xlsx(file_path: str, sheet_name: str | None = None) -> list[dict]:
    """ Читает xlsx файл и возвращает список словарей """
    wb = load_workbook(filename=file_path, read_only=True, data_only=True)

    if sheet_name is None:
        ws = wb.active
    else:
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"Лист '{sheet_name}' не найден. Доступные листы: {wb.sheetnames}")
        ws = wb[sheet_name]

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []  # пустой файл

    # Первая строка — заголовки
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]

    result = []
    for row in rows[1:]:
        # Если в строке меньше значений, чем заголовков - дополняем None
        row_values = list(row) + [None] * (len(headers) - len(row))
        row_dict = dict(zip(headers, row_values))
        result.append(row_dict)

    return result
